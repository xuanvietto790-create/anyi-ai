"""Local attachment storage. Files and extracted text stay on E: by default."""
import asyncio
import base64
import io
import json
import os
import uuid
import warnings
import zipfile
import threading
from pathlib import Path

from fastapi import HTTPException, UploadFile
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from docx import Document

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_TEXT = 1_000_000
IMAGE_TYPES = {'PNG': 'image/png', 'JPEG': 'image/jpeg', 'WEBP': 'image/webp', 'GIF': 'image/gif'}
claim_lock = threading.Lock()


def storage_dir():
    path = Path(os.getenv('ANYI_DATA_DIR', 'E:/Codex/work/anyi-ai-data')) / 'attachments'
    path.mkdir(parents=True, exist_ok=True)
    return path


def read_attachment(identifier: str, owner=None):
    try:
        if str(uuid.UUID(identifier)) != identifier:
            raise ValueError()
    except ValueError:
        raise HTTPException(400, '附件编号无效') from None
    try:
        metadata = json.loads((storage_dir() / f'{identifier}.json').read_text(encoding='utf-8'))
        if owner is not None and metadata.get('owner') != owner:
            raise HTTPException(404, '附件不存在或不属于当前账号')
        return metadata
    except (FileNotFoundError, ValueError):
        raise HTTPException(404, '附件已丢失，请重新上传') from None


def attachment_bytes(identifier: str):
    try:
        return (storage_dir() / f'{identifier}.bin').read_bytes()
    except FileNotFoundError:
        raise HTTPException(404, '附件原件已丢失，请重新上传') from None


def parse_file(data: bytes, name: str, progress=None):
    suffix = Path(name).suffix.lower()
    ocr_pages = 0
    if suffix in {'.png', '.jpg', '.jpeg', '.webp', '.gif'}:
        try:
            with warnings.catch_warnings():
                warnings.simplefilter('error', Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(data)) as image:
                    if image.format not in IMAGE_TYPES or max(image.size) > 8192:
                        raise ValueError()
                    media_type = IMAGE_TYPES[image.format]
                    image.verify()
            return {'kind': 'image', 'media_type': media_type, 'text': '', 'truncated': False}
        except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError, Image.DecompressionBombWarning):
            raise HTTPException(400, '图片损坏、格式不支持或尺寸过大，请上传有效图片（最长边 8192 像素）') from None
    try:
        if suffix == '.pdf':
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(400, 'PDF 已加密，请解密后再上传')
            if len(reader.pages) > 1000:
                raise HTTPException(400, 'PDF 最多支持 1000 页，请拆分后上传')
            parts, count, has_text = [], 0, False
            missing_pages = []
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text() or ''
                if len(''.join(extracted.split())) < 30:
                    missing_pages.append(i)
                has_text = has_text or bool(extracted.strip())
                value = f'\n[第 {i + 1} 页]\n{extracted}'
                parts.append(value)
                count += len(value)
                if progress:
                    progress(int(40 * (i + 1) / len(reader.pages)), f'读取页面 {i + 1}/{len(reader.pages)}')
                if count > MAX_TEXT:
                    break
            if missing_pages:
                from pdf_ocr import recognize_pages
                if progress:
                    progress(40, f'准备文字识别，共 {len(missing_pages)} 页')
                recognized = recognize_pages(data, missing_pages, progress) if progress else recognize_pages(data, missing_pages)
                ocr_pages = len(missing_pages)
                for i, value in recognized.items():
                    original = parts[i].split('\n', 2)[-1]
                    if len(value.strip()) > len(original.strip()):
                        parts[i] = f'\n[第 {i + 1} 页，OCR 识别，可能有误字]\n{value}'
                        has_text = True
            text = '\n'.join(parts)
            if not has_text:
                raise HTTPException(400, '此 PDF 未识别到文字，请确认页面不是空白，或上传更清晰的文件')
        elif suffix == '.docx':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 40 * 1024 * 1024 or len(archive.infolist()) > 2000:
                    raise HTTPException(400, 'Word 文件解压后过大，请精简后上传')
            document = Document(io.BytesIO(data))
            text = '\n'.join(p.text for p in document.paragraphs)
            for table in document.tables:
                text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
        elif suffix in {'.txt', '.md', '.csv', '.json', '.py', '.js', '.ts', '.html', '.css', '.log'}:
            try:
                text = data.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = data.decode('gb18030')
            if '\x00' in text:
                raise HTTPException(400, '文件不是可读取的文本')
        else:
            raise HTTPException(415, '支持图片、PDF、DOCX、TXT、Markdown、CSV 和常见代码文件；旧版 DOC 请另存为 DOCX')
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, '文件无法解析，请确认文件完整、格式正确且未加密') from None
    if not text.strip():
        raise HTTPException(400, '文件没有可读取的文字')
    return {'kind': 'document', 'media_type': 'application/octet-stream', 'text': text[:MAX_TEXT], 'truncated': len(text) > MAX_TEXT, 'ocr_pages': ocr_pages}


async def store_upload(file: UploadFile, owner, progress=None):
    try:
        data = await file.read(MAX_FILE_SIZE + 1)
    finally:
        await file.close()
    if not data or len(data) > MAX_FILE_SIZE:
        raise HTTPException(413, '文件不能为空，单个文件最大 10 MB')
    name = (file.filename or '附件').replace('\\', '/').rsplit('/', 1)[-1][:160]
    parsed = await asyncio.to_thread(parse_file, data, name, progress)
    if progress:
        progress(98, '保存解析结果')
    identifier = str(uuid.uuid4())
    metadata = {'id': identifier, 'name': name, 'size': len(data), 'text_limit': MAX_TEXT, 'owner': owner, **parsed}
    folder = storage_dir()
    if sum(p.stat().st_size for p in folder.glob('*.bin')) + len(data) > 500 * 1024 * 1024:
        raise HTTPException(507, '本地附件已达到 500 MB，请清理附件目录后重试')
    (folder / f'{identifier}.bin').write_bytes(data)
    (folder / f'{identifier}.json').write_text(json.dumps(metadata, ensure_ascii=False), encoding='utf-8')
    return {k: v for k, v in metadata.items() if k not in {'text', 'media_type', 'owner'}}


def claim_legacy(identifiers, owner):
    with claim_lock:
        attachments = [read_attachment(identifier) for identifier in set(identifiers)]
        if any(a.get('owner') not in {None, owner} for a in attachments):
            raise HTTPException(409, '部分旧附件已归属于其他账号，不能导入，请重新上传所需文件')
        for attachment in attachments:
            if attachment.get('owner') == owner:
                continue
            attachment['owner'] = owner
            path = storage_dir() / f'{attachment["id"]}.json'
            temporary = path.with_suffix('.json.tmp')
            temporary.write_text(json.dumps(attachment, ensure_ascii=False), encoding='utf-8')
            temporary.replace(path)
    return len(attachments)


def image_block(attachment):
    return {'type': 'image', 'source': {'type': 'base64', 'media_type': attachment['media_type'], 'data': base64.b64encode(attachment_bytes(attachment['id'])).decode()}}
