import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import httpx
from docx import Document
from PIL import Image
from pypdf import PdfWriter

import main
from auth_service import protect_api
from attachments import MAX_FILE_SIZE, MAX_TEXT

REAL_CLIENT = httpx.AsyncClient


def stream_response(thinking=False, truncated=False):
    events = []
    if thinking:
        events.append({'type': 'content_block_delta', 'delta': {'type': 'thinking_delta', 'thinking': '推理内容'}})
    events += [
        {'type': 'content_block_delta', 'delta': {'type': 'text_delta', 'text': '回答 [1]'}},
        {'type': 'message_delta', 'delta': {'stop_reason': 'max_tokens' if truncated else 'end_turn'}},
        {'type': 'message_stop'},
    ]
    return httpx.Response(200, text=''.join('data: ' + json.dumps(e) + '\n\n' for e in events))


class AppTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        Path('E:/Codex/tmp').mkdir(parents=True, exist_ok=True)
        self.directory = tempfile.TemporaryDirectory(prefix='anyi-ai-test-', dir='E:/Codex/tmp')
        self.env = patch.dict(os.environ, {'ANYI_DATA_DIR': self.directory.name, 'DEEPSEEK_API_KEY': 'test-secret'})
        self.env.start()
        async def test_account(request: main.Request):
            request.state.account = {'id': 'test-user', 'scope': 'test-user', 'nickname': '测试用户'}
        main.app.dependency_overrides[protect_api] = test_account
        self.client = REAL_CLIENT(transport=httpx.ASGITransport(app=main.app), base_url='http://test')

    async def asyncTearDown(self):
        await self.client.aclose()
        main.app.dependency_overrides.pop(protect_api, None)
        self.env.stop()
        self.directory.cleanup()

    def upstream(self, handler):
        return patch.object(main.httpx, 'AsyncClient', lambda **kwargs: REAL_CLIENT(transport=httpx.MockTransport(handler), **kwargs))

    async def upload(self, name, data, mime='application/octet-stream'):
        return await self.client.post('/api/attachments', files={'file': (name, data, mime)})

    async def ask(self, **options):
        return await self.client.post('/api/chat', json={'messages': [{'role': 'user', 'content': '测试问题'}], **options})

    async def test_health_and_static(self):
        self.assertEqual((await self.client.get('/')).status_code, 200)
        self.assertIn('thinking', (await self.client.get('/api/health')).json()['features'])
        self.assertEqual((await self.client.get('/static/vendor/purify.js')).status_code, 200)

    async def test_attachment_progress_completion(self):
        import uuid
        identifier = str(uuid.uuid4())
        response = await self.client.post('/api/attachments?task_id=' + identifier, files={'file': ('notes.txt', b'hello progress', 'text/plain')})
        self.assertEqual(response.status_code, 200)
        progress = await self.client.get('/api/attachments/progress/' + identifier)
        self.assertEqual(progress.status_code, 200)
        self.assertEqual(progress.json(), {'percent': 100, 'stage': '解析完成'})

    async def test_invalid_history(self):
        for messages in [[], [{'role': 'assistant', 'content': 'wrong'}], [{'role': 'user', 'content': 'a'}, {'role': 'user', 'content': 'b'}]]:
            self.assertEqual((await self.ask(messages=messages)).status_code, 422)

    async def test_missing_key(self):
        with patch.dict(os.environ, {'DEEPSEEK_API_KEY': ''}):
            self.assertEqual((await self.ask()).status_code, 503)

    async def test_stream_thinking_and_truncation(self):
        def handler(request):
            payload = json.loads(request.content)
            self.assertEqual(payload['thinking']['type'], 'enabled')
            self.assertEqual(payload['model'], 'deepseek-v4-pro')
            self.assertTrue(str(request.url).endswith('/anthropic/v1/messages'))
            return stream_response(thinking=True, truncated=True)
        with self.upstream(handler):
            response = await self.ask(thinking=True)
        self.assertIn('event: thinking', response.text)
        self.assertIn('event: notice', response.text)
        self.assertIn('event: done', response.text)

    async def test_search_sources_and_context(self):
        def handler(request):
            if request.url.path.endswith('/search'):
                return httpx.Response(200, json={'results': [{'title': '真实来源', 'url': 'https://example.org/article', 'content': '测试资料'}, {'title': '无效链接', 'url': 'javascript:alert(1)'}]})
            self.assertIn('测试资料', json.loads(request.content)['messages'][-1]['content'])
            return stream_response()
        with self.upstream(handler):
            response = await self.ask(search=True)
        self.assertIn('https://example.org/article', response.text)
        self.assertNotIn('javascript:', response.text)

    async def test_search_failure_does_not_claim_search_success(self):
        with self.upstream(lambda request: httpx.Response(403, text='secret upstream detail')):
            response = await self.ask(search=True)
        self.assertIn('event: error', response.text)
        self.assertNotIn('event: done', response.text)
        self.assertNotIn('secret upstream detail', response.text)

    async def test_empty_search_is_disclosed(self):
        with self.upstream(lambda request: httpx.Response(200, json={'results': []}) if request.url.path.endswith('/search') else stream_response()):
            response = await self.ask(search=True)
        self.assertIn('event: notice', response.text)
        self.assertIn('event: done', response.text)

    async def test_documents_and_truncation(self):
        result = await self.upload('notes.txt', ('安' * (MAX_TEXT + 1)).encode())
        self.assertEqual(result.status_code, 200)
        self.assertTrue(result.json()['truncated'])
        self.assertNotIn('text', result.json())
        document = Document(); document.add_paragraph('Word document text')
        table = document.add_table(rows=1, cols=2); table.cell(0, 0).text = 'price'; table.cell(0, 1).text = '42'
        buf = io.BytesIO(); document.save(buf)
        word = await self.upload('notes.docx', buf.getvalue())
        self.assertEqual(word.status_code, 200)
        def handler(request):
            content = json.loads(request.content)['messages'][-1]['content']
            self.assertIn('Word document text', content); self.assertIn('42', content)
            return stream_response()
        with self.upstream(handler):
            response = await self.ask(messages=[{'role': 'user', 'content': '总结', 'attachments': [word.json()['id']]}])
        self.assertIn('event: done', response.text)

    async def test_scanned_pdf_and_invalid_files(self):
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100); buf = io.BytesIO(); writer.write(buf)
        self.assertEqual((await self.upload('scan.pdf', buf.getvalue())).status_code, 400)
        self.assertEqual((await self.upload('fake.png', b'not an image')).status_code, 400)
        self.assertEqual((await self.upload('old.doc', b'invalid')).status_code, 415)
        self.assertEqual((await self.upload('huge.txt', b'a' * (MAX_FILE_SIZE + 1))).status_code, 413)

    async def test_image_vision_cache_and_preview(self):
        buf = io.BytesIO(); Image.new('RGB', (20, 20), 'red').save(buf, format='PNG')
        image = (await self.upload('sample.png', buf.getvalue())).json()
        self.assertEqual((await self.client.get(f'/api/attachments/{image["id"]}/preview')).headers['content-type'], 'image/png')
        calls = []
        def handler(request):
            payload = json.loads(request.content); calls.append(payload['model'])
            if payload['model'] == 'deepseek-flash':
                self.assertEqual(payload['messages'][0]['content'][0]['type'], 'image')
                return httpx.Response(200, json={'content': [{'type': 'text', 'text': '红色图片'}]})
            self.assertIn('红色图片', payload['messages'][0]['content'])
            return stream_response()
        with self.upstream(handler):
            for _ in range(2):
                result = await self.ask(messages=[{'role': 'user', 'content': '分析图片', 'attachments': [image['id']]}])
                self.assertIn('event: done', result.text)
        self.assertEqual(calls.count('deepseek-flash'), 1)

    async def test_missing_attachments_and_auth_error(self):
        self.assertEqual((await self.ask(messages=[{'role': 'user', 'content': 'a', 'attachments': ['../../secret']}])).status_code, 400)
        with self.upstream(lambda request: httpx.Response(401, text='test-secret')):
            result = await self.ask()
        self.assertIn('event: error', result.text)
        self.assertNotIn('test-secret', result.text)

    async def test_old_context_limits_removed(self):
        messages = [{'role': 'user' if i % 2 == 0 else 'assistant', 'content': '长上下文' * 500} for i in range(200)]
        messages.append({'role': 'user', 'content': '安' * 30000})
        result = await self.client.post('/api/context', json={'messages': messages})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()['limit'], 1048576)
        self.assertTrue(result.json()['fits'])

    async def test_over_budget_stops_before_model_request(self):
        count = {'input_tokens': 1048577, 'limit': 1048576, 'reserved_output_tokens': 8192, 'fits': False}
        with patch.object(main, 'budget', return_value=count):
            result = await self.ask()
        self.assertIn('event: error', result.text)
        self.assertIn('1M', result.text)
        self.assertNotIn('event: done', result.text)

    async def test_voice_audio_validation(self):
        response = await self.client.post('/api/speech/transcribe', files={'file': ('empty.wav', b'', 'audio/wav')})
        self.assertEqual(response.status_code, 413)
        response = await self.client.post('/api/speech/transcribe', files={'file': ('bad.wav', b'not audio', 'audio/wav')})
        # Missing ASR model returns 503; a ready model rejects corrupt audio with 400.
        self.assertIn(response.status_code, [400, 503])
        self.assertEqual((await self.client.post('/api/speech/synthesize', json={'text': '', 'rate': 1})).status_code, 422)

    async def test_history_roundtrip_and_validation(self):
        chats = [{'id': 'c1', 'title': '测试对话', 'createdAt': 1000, 'updatedAt': 2000,
                  'messages': [{'role': 'user', 'content': '你好'}, {'role': 'assistant', 'content': '你好！'}]}]
        self.assertEqual((await self.client.put('/api/history', json={'chats': chats})).status_code, 200)
        self.assertEqual((await self.client.get('/api/history')).json()['chats'], chats)
        # 全量替换：旧数据被清掉
        self.assertEqual((await self.client.put('/api/history', json={'chats': []})).status_code, 200)
        self.assertEqual((await self.client.get('/api/history')).json()['chats'], [])
        # 非法请求体
        self.assertEqual((await self.client.put('/api/history', json={'chats': 'nope'})).status_code, 422)


if __name__ == '__main__':
    unittest.main()
